"""Editable PowerPoint and offline print HTML from one reviewed French source.
Developer dependency: python-pptx. Screenshots contain synthetic demo data only.
"""
from pathlib import Path
import base64, html, json, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from PIL import Image

root=Path(__file__).resolve().parents[1]
docs=root/'docs'
slides=json.loads((docs/'guide-utilisation/contenu.json').read_text())
shots=docs/'captures/guide-utilisation'
navy='111D33'; blue='1759B5'; burgundy='892C4D'; muted='596880'; soft='F4F6FA'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
prs.core_properties.title='CHKPT — Guide d’utilisation complet'
prs.core_properties.subject='Prototype au 30 octobre 2026 · captures fictives · parcours et confidentialité'
prs.core_properties.author='CHKPT by RED LOOP'

def text(slide,x,y,w,h,content,size,color=navy,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left=tf.margin_right=Inches(.02);tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content if isinstance(content,list) else [content]):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(11 if isinstance(content,list) else 0);p.line_spacing=1.15
    return box

body=[]
for i,entry in enumerate(slides):
    image=shots/(entry['image']+'.png')
    if not image.exists():raise FileNotFoundError(image)
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    bg=slide.background.fill;bg.solid();bg.fore_color.rgb=RGBColor.from_string('FFFFFF')
    bar=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,Inches(.12));bar.fill.solid();bar.fill.fore_color.rgb=RGBColor.from_string(blue);bar.line.fill.background()
    text(slide,.65,.39,8,.25,'CHKPT  /  '+entry['chapter'],10,blue,True)
    text(slide,.65,.95,8.0,1.1,entry['title'],30,navy,True)
    text(slide,.72,2.2,7.55,3.58,[f'{n+1}.  {step}' for n,step in enumerate(entry['steps'])],17)
    note=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(.65),Inches(6.05),Inches(7.85),Inches(.79));note.fill.solid();note.fill.fore_color.rgb=RGBColor.from_string(soft);note.line.fill.background()
    text(slide,.85,6.18,7.45,.58,entry['note'],11.7,burgundy)
    panel=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(9.0),Inches(.48),Inches(3.55),Inches(6.45));panel.fill.solid();panel.fill.fore_color.rgb=RGBColor.from_string(soft);panel.line.fill.background()
    with Image.open(image) as pic:ratio=pic.width/pic.height
    height=6.05;width=height*ratio
    slide.shapes.add_picture(str(image),Inches(9.0+(3.55-width)/2),Inches(.67),width=Inches(width),height=Inches(height))
    text(slide,.65,7.13,10,.18,'30 OCT. 2026 · Données fictives · Guide du prototype · Captures réelles de la version testée',8,muted)
    text(slide,11.65,7.1,1,.23,f'{i+1:02} / {len(slides):02}',10,blue,True)
    slide.notes_slide.notes_text_frame.text=entry['title']+'\n\n'+'\n'.join(entry['steps'])+'\n\n'+entry['note']
    encoded=base64.b64encode(image.read_bytes()).decode()
    steps=''.join('<li>'+html.escape(step)+'</li>' for step in entry['steps'])
    body.append(f'''<section><header><b>CHKPT</b><span>{html.escape(entry['chapter'])}</span></header><main><div class="copy"><h1>{html.escape(entry['title'])}</h1><ol>{steps}</ol><aside>{html.escape(entry['note'])}</aside></div><figure><img src="data:image/png;base64,{encoded}" alt="Capture du prototype : {html.escape(entry['title'])}"><figcaption>Capture réelle · données fictives</figcaption></figure></main><footer><span>30 octobre 2026 · Guide du prototype · Aucun logiciel de programmation nécessaire pour le lire</span><strong>{i+1:02} / {len(slides):02}</strong></footer></section>''')
prs.save(docs/'GUIDE_UTILISATION_COMPLET.pptx')
css='''@page{size:A4 landscape;margin:0}*{box-sizing:border-box}body{margin:0;background:#fff;color:#111d33;font-family:Arial,sans-serif}section{width:297mm;height:210mm;padding:10mm 14mm;border-top:2mm solid #1759b5;position:relative;break-after:page}header{display:flex;justify-content:space-between;color:#1759b5;font-size:10px;letter-spacing:1px}header b{font-size:18px;letter-spacing:2px}main{height:174mm;display:grid;grid-template-columns:1fr 67mm;gap:13mm;align-items:center}.copy{min-width:0}h1{font-size:29px;line-height:1.15;letter-spacing:-.6px;margin:0 0 8mm}ol{margin:0;padding-left:7mm}li{font-size:16px;line-height:1.38;margin:0 0 4mm;color:#34435a}aside{margin-top:7mm;padding:4mm 5mm;background:#f4f6fa;border-left:1mm solid #892c4d;font-size:12px;line-height:1.4;color:#892c4d}figure{margin:0;text-align:center;background:#f4f6fa;border-radius:5mm;padding:3mm}img{width:100%;height:153mm;object-fit:contain}figcaption{font-size:8px;color:#596880;margin-top:2mm}footer{position:absolute;bottom:6mm;left:14mm;right:14mm;display:flex;justify-content:space-between;font-size:9px;color:#596880}footer strong{color:#1759b5}'''
if len(sys.argv)>1:
    Path(sys.argv[1]).write_text('<!doctype html><html lang="fr"><meta charset="utf-8"><title>CHKPT — Guide d’utilisation complet</title><style>'+css+'</style><body>'+''.join(body)+'</body></html>')
print(f'PowerPoint créé : {len(slides)} diapositives modifiables.')
